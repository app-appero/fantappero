"""Integration tests for the per-league bulk Excel roster export/import."""

from __future__ import annotations

import io
import re
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from tests.integration.database.helpers import create_engine_for_url

from database.enums import FantasyRole, LeagueMemberRole
from database.session import create_session_factory
from fantasy_teams.models import FantasyRosterSlot
from leagues.models.competition import Competition
from leagues.models.league_membership import LeagueMembership
from mail.capture import get_captured_emails
from sports_data.listone.models import RoleAssignment
from sports_data.roster.models import Athlete


def _register_and_login(client: TestClient, email: str) -> tuple[str, UUID]:
    client.post(
        "/auth/register",
        json={"email": email, "password": "Password123!", "displayName": email.split("@")[0]},
    )
    match = re.search(r"token=([A-Za-z0-9_-]+)", get_captured_emails()[-1].message.text_body)
    assert match
    client.post("/auth/verify-email", json={"token": match.group(1)})
    login = client.post(
        "/auth/login",
        json={"email": email, "password": "Password123!"},
    )
    assert login.status_code == 200
    return login.json()["accessToken"], UUID(login.json()["user"]["id"])


@pytest.fixture
def db_session(db_url: str) -> Session:
    engine = create_engine_for_url(db_url)
    session = create_session_factory(engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture
def competition_ids(db_session: Session) -> list[str]:
    rows = db_session.scalars(select(Competition).order_by(Competition.name.asc())).all()
    assert len(rows) >= 3
    return [str(row.id) for row in rows[:3]]


def _create_league(
    client: TestClient,
    token: str,
    competition_ids: list[str],
    name: str,
) -> str:
    response = client.post(
        "/leagues",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "name": name,
            "seasonYear": 2026,
            "competitionIds": competition_ids,
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


def _seed_athlete(
    db_session: Session,
    provider_id: int,
    name: str,
    *,
    role: FantasyRole = FantasyRole.A,
    season_year: int = 2026,
) -> Athlete:
    athlete = Athlete(provider_id=provider_id, canonical_name=name)
    db_session.add(athlete)
    db_session.flush()
    db_session.add(
        RoleAssignment(
            athlete_id=athlete.id,
            season_year=season_year,
            role=role,
            mapping_version="v1.0.0",
            provider_position_raw=role.value,
        )
    )
    db_session.commit()
    db_session.refresh(athlete)
    return athlete


def _sheet_names(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [sheet.title for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _bulk_xlsx(team_rows: dict[str, list[tuple[str, int]]]) -> bytes:
    workbook = Workbook()
    first = True
    for sheet_name, rows in team_rows.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = sheet_name
        sheet.append(["Calciatore", "Crediti"])
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_bulk_export_has_one_sheet_per_coach(
    client: TestClient,
    db_session: Session,
    competition_ids: list[str],
) -> None:
    owner_token, owner_id = _register_and_login(client, "bulk.owner@example.com")
    member_token, member_id = _register_and_login(client, "bulk.member@example.com")
    league_id = _create_league(client, owner_token, competition_ids, "Lega Bulk Export")
    db_session.add(
        LeagueMembership(league_id=UUID(league_id), user_id=member_id, role=LeagueMemberRole.MEMBER)
    )
    db_session.commit()

    owner_rosa = client.get(
        f"/leagues/{league_id}/rosa", headers={"Authorization": f"Bearer {owner_token}"}
    )
    member_rosa = client.get(
        f"/leagues/{league_id}/rosa", headers={"Authorization": f"Bearer {member_token}"}
    )
    assert owner_rosa.status_code == 200
    assert member_rosa.status_code == 200

    export = client.get(
        f"/leagues/{league_id}/amministrazione/rose/excel-tutte",
        headers={"Authorization": f"Bearer {owner_token}"},
    )
    assert export.status_code == 200
    assert "spreadsheetml" in export.headers["content-type"]
    names = _sheet_names(export.content)
    assert len(names) == 2
    assert {"bulk.owner", "bulk.member"} <= {n for n in names}

    forbidden = client.get(
        f"/leagues/{league_id}/amministrazione/rose/excel-tutte",
        headers={"Authorization": f"Bearer {member_token}"},
    )
    assert forbidden.status_code in {403, 401}


def test_bulk_import_assigns_each_coach_sheet_and_skips_unmatched(
    client: TestClient,
    db_session: Session,
    competition_ids: list[str],
) -> None:
    owner_token, owner_id = _register_and_login(client, "bulkimp.owner@example.com")
    member_token, member_id = _register_and_login(client, "bulkimp.member@example.com")
    league_id = _create_league(client, owner_token, competition_ids, "Lega Bulk Import")
    db_session.add(
        LeagueMembership(league_id=UUID(league_id), user_id=member_id, role=LeagueMemberRole.MEMBER)
    )
    db_session.commit()

    owner_rosa = client.get(
        f"/leagues/{league_id}/rosa", headers={"Authorization": f"Bearer {owner_token}"}
    )
    member_rosa = client.get(
        f"/leagues/{league_id}/rosa", headers={"Authorization": f"Bearer {member_token}"}
    )
    owner_team_id = owner_rosa.json()["id"]
    member_team_id = member_rosa.json()["id"]

    a1 = _seed_athlete(db_session, 9301, "Bulk Player One")
    a2 = _seed_athlete(db_session, 9302, "Bulk Player Two")

    data = _bulk_xlsx(
        {
            "bulkimp.owner": [(a1.canonical_name, 20)],
            "bulkimp.member": [(a2.canonical_name, 30)],
            "Foglio Sconosciuto": [("Nessuno", 1)],
        }
    )
    preview = client.post(
        f"/leagues/{league_id}/amministrazione/import-csv/anteprima-massiva",
        headers={"Authorization": f"Bearer {owner_token}"},
        files={"file": ("rose.xlsx", data, "application/octet-stream")},
    )
    assert preview.status_code == 200
    body = preview.json()
    assert body["canConfirm"] is True
    assert body["errorCount"] == 0
    assert body["skippedSheets"] == ["Foglio Sconosciuto"]
    assert {row["squadra"] for row in body["rows"]} == {
        owner_rosa.json()["name"],
        member_rosa.json()["name"],
    }
    import_id = body["importId"]

    confirm = client.post(
        f"/leagues/{league_id}/amministrazione/import-csv/{import_id}/conferma",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"resolutions": []},
    )
    assert confirm.status_code == 200
    assert confirm.json()["assignedCount"] == 2
    assert confirm.json()["preview"]["skippedSheets"] == ["Foglio Sconosciuto"]

    owner_slot = db_session.scalar(
        select(FantasyRosterSlot).where(
            FantasyRosterSlot.fantasy_team_id == UUID(owner_team_id),
            FantasyRosterSlot.athlete_id == a1.id,
        )
    )
    member_slot = db_session.scalar(
        select(FantasyRosterSlot).where(
            FantasyRosterSlot.fantasy_team_id == UUID(member_team_id),
            FantasyRosterSlot.athlete_id == a2.id,
        )
    )
    assert owner_slot is not None
    assert member_slot is not None


def test_bulk_import_rejects_csv_file(
    client: TestClient,
    db_session: Session,
    competition_ids: list[str],
) -> None:
    owner_token, _ = _register_and_login(client, "bulkcsv.owner@example.com")
    league_id = _create_league(client, owner_token, competition_ids, "Lega Bulk CSV Rejected")
    client.get(f"/leagues/{league_id}/rosa", headers={"Authorization": f"Bearer {owner_token}"})

    response = client.post(
        f"/leagues/{league_id}/amministrazione/import-csv/anteprima-massiva",
        headers={"Authorization": f"Bearer {owner_token}"},
        files={
            "file": (
                "rose.csv",
                b"squadra,provider_id,nome,crediti\nA,1,B,1\n",
                "text/csv",
            )
        },
    )
    assert response.status_code == 400
    assert response.json()["code"] == "xlsx_required"
