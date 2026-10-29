"""Unit tests for the two-column Excel roster file."""

from __future__ import annotations

import io
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from auth.exceptions import ValidationAuthError
from fantasy_teams.csv_import import PreviewRow, PreviewRowIssue
from fantasy_teams.csv_preview import count_preview_issues, skip_unmatched_excel_rows
from fantasy_teams.xlsx_roster import (
    XLSX_HEADERS,
    is_xlsx,
    parse_bulk_roster_xlsx,
    parse_roster_xlsx,
    rows_to_bulk_xlsx_bytes,
    rows_to_xlsx_bytes,
    safe_sheet_name,
    template_xlsx_bytes,
)


def _sheet_names(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [sheet.title for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _sheet_rows(data: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        assert sheet is not None
        return [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def test_template_has_only_player_and_credits() -> None:
    data = template_xlsx_bytes()
    assert is_xlsx(data)
    rows = _sheet_rows(data)
    assert rows == [("Calciatore", "Crediti"), ("Nome Calciatore", 10)]
    parsed = parse_roster_xlsx(data, team_name="Rosa demo")
    assert len(parsed) == 1
    assert parsed[0].squadra == "Rosa demo"
    assert parsed[0].provider_id is None
    assert parsed[0].nome == "Nome Calciatore"
    assert parsed[0].crediti == 10


def test_export_round_trip_ignores_team_name_in_the_file() -> None:
    data = rows_to_xlsx_bytes([("L. Martinez", 50), ("N. Barella", 35)])
    assert _sheet_rows(data)[0] == XLSX_HEADERS
    parsed = parse_roster_xlsx(data, team_name="Allenatore IA 01")
    assert [(row.nome, row.crediti, row.provider_id, row.squadra) for row in parsed] == [
        ("L. Martinez", 50, None, "Allenatore IA 01"),
        ("N. Barella", 35, None, "Allenatore IA 01"),
    ]


def test_blank_excel_rows_are_skipped() -> None:
    data = rows_to_xlsx_bytes([("L. Martinez", 10), ("", None)])
    parsed = parse_roster_xlsx(data, team_name="Rosa demo")
    assert len(parsed) == 1
    assert parsed[0].nome == "L. Martinez"


def test_header_without_player_column_is_rejected() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Crediti"])
    sheet.append([10])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ValidationAuthError) as caught:
        parse_roster_xlsx(buffer.getvalue(), team_name="Rosa demo")
    assert caught.value.code == "xlsx_invalid_header"


def test_unknown_player_row_becomes_a_warning_skip() -> None:
    row = PreviewRow(
        row_number=2,
        squadra="Rosa demo",
        provider_id=None,
        nome="Sconosciuto",
        crediti=8,
        status="error",
        fantasy_team_id=str(uuid4()),
        issues=[
            PreviewRowIssue(
                code="athlete_not_found",
                message="Nessun calciatore trovato per il nome «Sconosciuto».",
                severity="error",
            )
        ],
    )
    matched = PreviewRow(
        row_number=3,
        squadra="Rosa demo",
        provider_id=None,
        nome="L. Martinez",
        crediti=50,
        status="error",
        fantasy_team_id=row.fantasy_team_id,
        athlete_id=str(uuid4()),
        issues=[
            PreviewRowIssue(
                code="athlete_already_on_team",
                message="Il calciatore è già nella rosa di questa squadra.",
                severity="error",
            )
        ],
    )
    skip_unmatched_excel_rows([row, matched])
    assert row.status == "skipped"
    assert row.athlete_id is None
    assert row.slot_index is None
    assert row.issues[0].severity == "warning"
    assert matched.status == "error"
    errors, warnings = count_preview_issues([row, matched])
    assert errors == 1
    assert warnings == 1


def test_safe_sheet_name_strips_forbidden_characters_and_truncates() -> None:
    used: set[str] = set()
    name = safe_sheet_name("Mario / Rossi [capitano]: *top*", used_lower=used)
    assert "/" not in name and "[" not in name and ":" not in name and "*" not in name
    assert len(name) <= 31


def test_safe_sheet_name_deduplicates_case_insensitively() -> None:
    used: set[str] = set()
    first = safe_sheet_name("Mario Rossi", used_lower=used)
    second = safe_sheet_name("mario rossi", used_lower=used)
    assert first == "Mario Rossi"
    assert second != first
    assert second.lower().startswith("mario rossi")


def test_safe_sheet_name_falls_back_when_blank() -> None:
    used: set[str] = set()
    assert safe_sheet_name("   ", used_lower=used) == "Squadra"


def test_bulk_export_round_trip_one_sheet_per_team() -> None:
    data = rows_to_bulk_xlsx_bytes(
        [
            ("Mario Rossi", [("L. Martinez", 50)]),
            ("Luigi Bianchi", [("N. Barella", 35)]),
        ]
    )
    assert _sheet_names(data) == ["Mario Rossi", "Luigi Bianchi"]
    parsed, skipped = parse_bulk_roster_xlsx(
        data,
        team_name_by_sheet_name_lower={
            "mario rossi": "Squadra Rossi",
            "luigi bianchi": "Squadra Bianchi",
        },
    )
    assert skipped == []
    assert [(row.squadra, row.nome, row.crediti) for row in parsed] == [
        ("Squadra Rossi", "L. Martinez", 50),
        ("Squadra Bianchi", "N. Barella", 35),
    ]


def test_bulk_export_without_teams_produces_valid_placeholder_workbook() -> None:
    data = rows_to_bulk_xlsx_bytes([])
    assert _sheet_names(data) == ["Nessuna squadra"]


def test_bulk_import_skips_unrecognized_sheets_silently() -> None:
    data = rows_to_bulk_xlsx_bytes(
        [
            ("Mario Rossi", [("L. Martinez", 50)]),
            ("Foglio Misterioso", [("N. Barella", 35)]),
        ]
    )
    parsed, skipped = parse_bulk_roster_xlsx(
        data, team_name_by_sheet_name_lower={"mario rossi": "Squadra Rossi"}
    )
    assert skipped == ["Foglio Misterioso"]
    assert [row.nome for row in parsed] == ["L. Martinez"]


def test_bulk_import_allows_empty_matched_sheet() -> None:
    data = rows_to_bulk_xlsx_bytes(
        [
            ("Mario Rossi", [("L. Martinez", 50)]),
            ("Luigi Bianchi", []),
        ]
    )
    parsed, skipped = parse_bulk_roster_xlsx(
        data,
        team_name_by_sheet_name_lower={
            "mario rossi": "Squadra Rossi",
            "luigi bianchi": "Squadra Bianchi",
        },
    )
    assert skipped == []
    assert [row.nome for row in parsed] == ["L. Martinez"]


def test_bulk_import_raises_when_no_sheet_matches() -> None:
    data = rows_to_bulk_xlsx_bytes([("Foglio Misterioso", [("N. Barella", 35)])])
    with pytest.raises(ValidationAuthError) as caught:
        parse_bulk_roster_xlsx(data, team_name_by_sheet_name_lower={"mario rossi": "Squadra"})
    assert caught.value.code == "xlsx_bulk_no_match"
