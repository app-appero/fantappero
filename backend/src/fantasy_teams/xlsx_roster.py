"""Excel roster template, export and import: player name and credits only."""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from auth.exceptions import ValidationAuthError
from fantasy_teams.csv_import import (
    CSV_MAX_BYTES,
    CSV_MAX_ROWS,
    ParsedCsvRow,
    normalize_header,
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSX_TEMPLATE_FILENAME = "fantappero-import-rosa.xlsx"
XLSX_EXPORT_FILENAME = "fantappero-rose.xlsx"
XLSX_EXPORT_ALL_FILENAME = "fantappero-rose-lega.xlsx"
XLSX_HEADERS = ("Calciatore", "Crediti")
#: Excel non ammette questi caratteri nel nome di un foglio, ed è lungo al
#: massimo 31 caratteri — qui per non duplicare la regola fra export e
#: import (devono concordare esattamente sullo stesso nome).
_SHEET_NAME_FORBIDDEN = re.compile(r"[\\/?*\[\]:]")
_SHEET_NAME_MAX_LENGTH = 31

_EXAMPLE_ROW = ("Nome Calciatore", 10)


def is_xlsx(data: bytes) -> bool:
    return data.startswith(b"PK\x03\x04")


def safe_sheet_name(name: str, *, used_lower: set[str]) -> str:
    """Nome foglio valido per Excel, univoco in ``used_lower`` (case-insensitive,
    Excel stesso non distingue maiuscole/minuscole nei nomi foglio).

    Export e import devono concordare esattamente su questa funzione: è
    quello che lega "il foglio si chiama come l'allenatore" a "il foglio
    che non combacia con nessun allenatore viene ignorato".
    """
    base = _SHEET_NAME_FORBIDDEN.sub(" ", name).strip() or "Squadra"
    base = base[:_SHEET_NAME_MAX_LENGTH]
    candidate = base
    suffix = 2
    while candidate.lower() in used_lower:
        tail = f" ({suffix})"
        candidate = base[: _SHEET_NAME_MAX_LENGTH - len(tail)] + tail
        suffix += 1
    used_lower.add(candidate.lower())
    return candidate


def rows_to_bulk_xlsx_bytes(team_rows: Sequence[tuple[str, list[tuple[object, ...]]]]) -> bytes:
    """Un foglio per squadra, nome foglio = nome allenatore (già reso univoco
    dal chiamante con ``safe_sheet_name``). ``team_rows`` è ``[(sheet_name,
    rows), ...]``, stesse due colonne Calciatore/Crediti di un export singolo.
    Una lega senza squadre produce un unico foglio segnaposto vuoto: un
    workbook senza fogli non è un .xlsx valido."""
    workbook = Workbook()
    first = True
    for sheet_name, rows in team_rows:
        sheet: Worksheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = sheet_name
        sheet.append(list(XLSX_HEADERS))
        for row in rows:
            sheet.append(list(row))
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 12
    if first:
        workbook.active.title = "Nessuna squadra"
        workbook.active.append(list(XLSX_HEADERS))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template_xlsx_bytes() -> bytes:
    return rows_to_xlsx_bytes([_EXAMPLE_ROW])


def rows_to_xlsx_bytes(rows: Iterable[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rosa"
    sheet.append(list(XLSX_HEADERS))
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 12
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parse_roster_xlsx(data: bytes, *, team_name: str) -> list[ParsedCsvRow]:
    """Read Calciatore + Crediti. The fantasy team stays out of the file."""
    if not data:
        raise ValidationAuthError("Il file Excel è vuoto.", code="csv_empty")
    if len(data) > CSV_MAX_BYTES:
        raise ValidationAuthError(
            "Il file Excel supera la dimensione massima consentita.",
            code="csv_too_large",
        )
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationAuthError(
            "Il file Excel non è leggibile. Usa il modello ufficiale .xlsx.",
            code="xlsx_invalid",
        ) from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise ValidationAuthError(
                "Il file Excel non contiene un foglio.",
                code="xlsx_invalid",
            )
        return _rows_from_sheet(sheet, team_name=team_name)
    finally:
        workbook.close()


def parse_bulk_roster_xlsx(
    data: bytes,
    *,
    team_name_by_sheet_name_lower: dict[str, str],
) -> tuple[list[ParsedCsvRow], list[str]]:
    """Un foglio per squadra: il nome del foglio DEVE combaciare (case-
    insensitive) con il nome allenatore di una squadra della lega — quelli
    già calcolati da ``safe_sheet_name`` all'export, passati qui come
    ``{nome_foglio.lower(): team.name}``. Un foglio che non combacia con
    nessuno viene ignorato in silenzio (nessun errore, nessuna riga): è la
    stessa idea del CSV con colonna `squadra` non riconosciuta, qui applicata
    all'intero foglio invece che riga per riga.

    Ritorna ``(righe_di_tutti_i_fogli_riconosciuti, nomi_fogli_ignorati)``:
    le righe hanno già `squadra` impostato al nome VERO della squadra
    (`team.name`), pronte per la stessa pipeline di preview/conferma già
    usata dall'import CSV multi-squadra (colonna `squadra`).
    """
    if not data:
        raise ValidationAuthError("Il file Excel è vuoto.", code="csv_empty")
    if len(data) > CSV_MAX_BYTES:
        raise ValidationAuthError(
            "Il file Excel supera la dimensione massima consentita.",
            code="csv_too_large",
        )
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationAuthError(
            "Il file Excel non è leggibile. Usa il modello ufficiale .xlsx.",
            code="xlsx_invalid",
        ) from exc

    rows: list[ParsedCsvRow] = []
    skipped: list[str] = []
    try:
        if not workbook.worksheets:
            raise ValidationAuthError(
                "Il file Excel non contiene un foglio.",
                code="xlsx_invalid",
            )
        for sheet in workbook.worksheets:
            team_name = team_name_by_sheet_name_lower.get((sheet.title or "").strip().lower())
            if team_name is None:
                skipped.append(sheet.title)
                continue
            rows.extend(_rows_from_sheet(sheet, team_name=team_name, allow_empty=True))
    finally:
        workbook.close()

    if not rows:
        raise ValidationAuthError(
            "Nessun foglio corrisponde al nome di un allenatore di questa lega.",
            code="xlsx_bulk_no_match",
        )
    return rows, skipped


def _rows_from_sheet(
    sheet: object, *, team_name: str, allow_empty: bool = False
) -> list[ParsedCsvRow]:
    name_index: int | None = None
    credit_index: int | None = None
    parsed: list[ParsedCsvRow] = []
    data_rows = 0
    sheet_rows = sheet.iter_rows(values_only=True)  # type: ignore[attr-defined]
    for row_number, raw in enumerate(sheet_rows, start=1):
        cells = tuple(raw)
        if all(_is_blank(cell) for cell in cells):
            continue
        if name_index is None or credit_index is None:
            mapped = [normalize_header(_cell_text(cell)) for cell in cells]
            if "nome" not in mapped or "crediti" not in mapped:
                raise ValidationAuthError(
                    "Il file Excel deve avere le colonne Calciatore e Crediti.",
                    code="xlsx_invalid_header",
                )
            name_index = mapped.index("nome")
            credit_index = mapped.index("crediti")
            continue

        data_rows += 1
        if data_rows > CSV_MAX_ROWS:
            raise ValidationAuthError(
                f"Il file Excel supera il massimo di {CSV_MAX_ROWS} righe dati.",
                code="csv_too_many_rows",
            )
        nome = _cell_text(cells[name_index]) if name_index < len(cells) else ""
        credit_cell = cells[credit_index] if credit_index < len(cells) else None
        crediti, invalid = _parse_credits(credit_cell)
        if not nome and crediti is None and not invalid:
            continue
        raw_map = {"nome": nome, "crediti": _cell_text(credit_cell)}
        if invalid:
            raw_map["_crediti_invalid"] = raw_map["crediti"]
        parsed.append(
            ParsedCsvRow(
                row_number=row_number,
                squadra=team_name,
                provider_id=None,
                nome=nome,
                crediti=crediti,
                raw=raw_map,
            )
        )

    if name_index is None:
        if allow_empty:
            return []
        raise ValidationAuthError("Il file Excel è vuoto.", code="csv_empty")
    if not parsed and not allow_empty:
        raise ValidationAuthError(
            "Il file Excel non contiene righe dati da importare.",
            code="csv_no_data_rows",
        )
    return parsed


def _parse_credits(value: object) -> tuple[int | None, bool]:
    if _is_blank(value):
        return None, False
    if isinstance(value, bool):
        return None, True
    if isinstance(value, int):
        return value, False
    if isinstance(value, float):
        if value.is_integer():
            return int(value), False
        return None, True
    text = _cell_text(value)
    try:
        return int(text), False
    except ValueError:
        return None, True


def _is_blank(value: object) -> bool:
    return value is None or str(value).strip() == ""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()
